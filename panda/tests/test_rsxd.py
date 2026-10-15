# =============================================================================
# RSXD SCAN LOGGER  --  acceptance tests
# =============================================================================
# Drives the REAL tkinter app with REAL key events (event_generate), then
# inspects the files it actually wrote. Not a mock in sight.
#
# Run:  python tests/test_rsxd.py
#
# Every popup entry point is monkeypatched to raise, so any test that would
# have produced a dialog fails loudly instead of silently passing.
# =============================================================================

import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
from pathlib import Path

APP_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(APP_DIR))

import tkinter as tk
from tkinter import filedialog, messagebox

import openpyxl

import rsxd_logger
from rsxd_core import CSV_COLUMNS, Config, Entry, Session

Entry0 = Entry(scan=-1)          # stand-in so "cell is empty" reads as scan -1

REAL_DAT_DIR = r"C:\Users\fng525\OneDrive - University of York\YORK 2025\LAB WORK\XMCD\data\dls\mm43464"
REAL_DAT_LO, REAL_DAT_HI = 64443, 64690

RESULTS = []


# =============================================================================
# HARNESS
# =============================================================================

class PopupDetected(AssertionError):
    pass


def _no_popups():
    def boom(*a, **k):
        raise PopupDetected(f"a modal dialog was raised: {a!r}")
    for mod in (messagebox, filedialog):
        for name in dir(mod):
            if name.startswith(("show", "ask")):
                setattr(mod, name, boom)


_no_popups()


class Harness:
    """One app instance in a private scratch directory."""

    def __init__(self, workdir: Path, config_text: str | None = None,
                 reuse_config: bool = False):
        self.workdir = workdir
        workdir.mkdir(parents=True, exist_ok=True)
        self.config_path = workdir / "rsxd_config.toml"
        if not reuse_config:
            text = config_text if config_text is not None else \
                (APP_DIR / "rsxd_config.toml").read_text(encoding="utf-8")
            repl = f"output_dir = '{workdir / 'output'}'"
            text = re.sub(r"^output_dir\s*=.*$", lambda _m: repl,
                          text, count=1, flags=re.M)
            self.config_path.write_text(text, encoding="utf-8")
        self.root = None
        self.app = None

    def start(self):
        self.root = tk.Tk()
        self.root.geometry("1300x840")
        self.app = rsxd_logger.RSXDLoggerFrame(self.root, config_path=self.config_path)
        self.app.pack(fill="both", expand=True)
        self.app.install_shortcuts(self.root)
        # Successive Tk roots in one process do not inherit OS keyboard focus.
        # Without this, event_generate would go nowhere and tests would pass
        # vacuously. Take focus explicitly and assert on it below.
        self.root.deiconify()
        self.root.lift()
        self.root.focus_force()
        self.pump(8)
        return self.app

    def pump(self, n=3):
        for _ in range(n):
            self.root.update_idletasks()
            self.root.update()

    # ---- real keyboard ------------------------------------------------------
    def cell_focus(self, app, key=None):
        """
        Put real keyboard focus on a grid cell.

        The app uses focus_set(), which is right on a normal desktop -- test 1
        run on its own proves it, because there the window does get OS focus.
        But this suite creates and destroys several Tk roots inside one
        process, and Windows does not grant OS focus to the later ones, so
        focus_set() has no window to attach to. focus_force() on the widget
        itself works regardless. Only the FIRST focus per app instance needs
        this; once the toplevel holds focus, the app's own Tab navigation
        drives everything from there.
        """
        if key is None:
            key = next((k for k in app._order() if not app.cell_vars[k].get()),
                       app._order()[0])
        app.cells[key].focus_force()
        self.pump(3)
        w = self.root.focus_get()
        assert w is app.cells[key], f"focus went to {w} instead of cell {key}"
        return key

    def focused(self):
        w = self.root.focus_get()
        assert w is not None, "no widget has focus — keystrokes would go nowhere"
        return w

    def type_text(self, text):
        """Send one <KeyPress> per character to whatever currently has focus."""
        w = self.focused()
        assert isinstance(w, tk.Entry), f"focus is on {w.winfo_class()}, not an Entry"
        for ch in str(text):
            w.event_generate("<KeyPress>", keysym=ch)
        self.pump(1)

    def key(self, keysym):
        self.focused().event_generate(f"<{keysym}>")
        self.pump(1)

    def clear_focused(self):
        w = self.root.focus_get()
        if isinstance(w, tk.Entry):
            w.select_range(0, "end")
            w.event_generate("<KeyPress>", keysym="BackSpace")
        self.pump(1)

    def log_text(self):
        return self.app.log_text.get("1.0", "end")

    def stop(self, clean=True):
        if self.root is not None:
            if clean:
                self.app.on_close()
            self.root.destroy()
            self.root = None


def check(name, condition, detail=""):
    RESULTS.append((name, bool(condition), detail))
    mark = "PASS" if condition else "FAIL"
    print(f"  [{mark}] {name}" + (f"   {detail}" if detail else ""))
    return bool(condition)


def banner(n, title):
    print(f"\n{'=' * 78}\nTEST {n}  {title}\n{'=' * 78}")


def scratch(name) -> Path:
    d = Path(tempfile.gettempdir()) / "rsxd_tests" / name
    if d.exists():
        shutil.rmtree(d, ignore_errors=True)
    d.mkdir(parents=True, exist_ok=True)
    return d


# =============================================================================
# TEST 1  --  full 4 x 13 grid, keyboard only
# =============================================================================
def test_1():
    banner(1, "Enter a full 4 x 13 grid using only the keyboard")
    h = Harness(scratch("t1"))
    app = h.start()
    cfg = app.cfg

    h.cell_focus(app)

    scan = 64700
    typed_order = []
    for _ in range(len(cfg.angles) * len(cfg.pol_labels)):
        focused = h.root.focus_get()
        key = next(k for k, w in app.cells.items() if w is focused)
        typed_order.append(key)
        h.type_text(str(scan))
        h.key("Tab")
        scan += 1

    expected_order = [(a, p) for a in cfg.angles for p in cfg.pol_labels]
    check("Tab visits all 52 cells", len(typed_order) == 52, f"{len(typed_order)} cells")
    check("Tab order = 4 channels at one angle, then next angle",
          typed_order == expected_order)

    filled = len(app.session.scans_for(app.session.ui_peak, app.session.ui_temp))
    check("all 52 cells stored", filled == 52, f"{filled}/52")
    check("header shows 0 empty", app.readouts["empty"].cget("text") == "0",
          f"EMPTY readout = {app.readouts['empty'].cget('text')}")

    spot = app.session.get(app.session.key(cfg.angles[0], cfg.pol_labels[0]))
    check("first cell = 64700", spot is not None and spot.scan == 64700,
          f"got {spot.scan if spot else None}")
    last = app.session.get(app.session.key(cfg.angles[-1], cfg.pol_labels[-1]))
    check("last cell = 64751", last is not None and last.scan == 64751,
          f"got {last.scan if last else None}")
    check("no duplicates flagged", not app.session.duplicates())

    app.save_current()
    h.pump(2)
    files = sorted(p.name for p in (h.workdir / "output").glob("*.xlsx"))
    check("one workbook written", files == ["CoTau_300K_64700-64751.xlsx"], str(files))
    h.stop()


# =============================================================================
# TEST 2  --  hard kill mid-entry, reopen, nothing lost
# =============================================================================
def test_2():
    banner(2, "Kill the app mid-entry. Reopen. Nothing is lost.")
    work = scratch("t2")
    victim = work / "victim.py"
    victim.write_text(f'''
import sys
sys.path.insert(0, r"{APP_DIR}")
sys.path.insert(0, r"{APP_DIR / "tests"}")
import os, tkinter as tk, rsxd_logger
root = tk.Tk(); root.geometry("1300x840")
app = rsxd_logger.RSXDLoggerFrame(root, config_path=r"{work / 'rsxd_config.toml'}")
app.pack(fill="both", expand=True)
for _ in range(6): root.update_idletasks(); root.update()
app.focus_first_empty()
for _ in range(3): root.update_idletasks(); root.update()
scan = 64810
for i in range(7):
    w = root.focus_get()
    for ch in str(scan): w.event_generate("<KeyPress>", keysym=ch)
    root.update_idletasks(); root.update()
    if i == 3:
        app.note_entry.focus_set()
        for ch in "beam dump": app.note_entry.event_generate("<KeyPress>", keysym=ch if ch!=" " else "space")
        root.update_idletasks(); root.update()
        app.leave_note()
    w = root.focus_get(); w.event_generate("<Tab>")
    root.update_idletasks(); root.update()
    scan += 1
print("VICTIM_ENTERED_7", flush=True)
# Hard kill: no atexit, no __del__, no WM_DELETE_WINDOW, no flush.
os._exit(1)
''', encoding="utf-8")

    Harness(work)            # writes the config
    proc = subprocess.run([sys.executable, str(victim)], capture_output=True, text=True,
                          timeout=180)
    check("victim process entered 7 scans then hard-exited",
          "VICTIM_ENTERED_7" in proc.stdout and proc.returncode == 1,
          f"rc={proc.returncode}")
    if proc.returncode != 1:
        print(proc.stdout[-2000:], proc.stderr[-2000:])

    sess_file = work / "output" / "rsxd_session.json"
    check("session file survived the kill", sess_file.is_file())

    h = Harness(work, reuse_config=True)
    app = h.start()
    scans = app.session.scans_for(app.session.ui_peak, app.session.ui_temp)
    check("all 7 scans restored", scans == list(range(64810, 64817)), str(scans))
    check("log says session restored", "Session restored: 7 scans" in h.log_text())

    notes = [e.note for e in app.session.entries.values() if e.note]
    check("per-cell note survived too", notes == ["beam dump"], str(notes))

    shown = [app.cell_vars[k].get() for k in app._order()[:7]]
    check("grid repainted with the restored values",
          shown == [str(n) for n in range(64810, 64817)], str(shown))
    h.stop()


# =============================================================================
# TEST 3  --  partial grid saves without complaint
# =============================================================================
def test_3():
    banner(3, "Save a grid with only six cells filled")
    h = Harness(scratch("t3"))
    app = h.start()
    h.cell_focus(app)
    for scan in range(64900, 64906):
        h.type_text(str(scan))
        h.key("Tab")

    before = len(h.log_text())
    app.save_current()          # would raise PopupDetected if it warned
    h.pump(2)
    new_log = h.log_text()[before:]

    out = h.workdir / "output"
    files = sorted(p.name for p in out.glob("*.xlsx"))
    check("saved with 6/52 cells, no dialog",
          files == ["CoTau_300K_64900-64905.xlsx"], str(files))
    check("log reports the partial fill plainly",
          "6/52 cells filled" in new_log,
          [l for l in new_log.splitlines() if "cells filled" in l])
    check("no warning words in the save log",
          not any(w in new_log.lower() for w in ("error", "incomplete", "cannot", "refus")))

    wb = openpyxl.load_workbook(out / files[0])
    ws = wb["Sheet1"]
    populated = sum(1 for r in range(3, 7) for c in range(3, 16)
                    if ws.cell(row=r, column=c).value is not None)
    check("workbook has exactly 6 scan cells, rest blank", populated == 6, str(populated))
    h.stop()


# =============================================================================
# TEST 4  --  duplicate scan number warns
# =============================================================================
def test_4():
    banner(4, "Enter a scan number that already exists elsewhere")
    h = Harness(scratch("t4"))
    app = h.start()

    h.cell_focus(app)
    h.type_text("64717")
    h.key("Tab")

    # Move to a DIFFERENT peak and temperature, then repeat the same number.
    app.peak_var.set("Fe-tau")
    app.temp_var.set("200")
    app.on_context_change()
    h.pump(2)

    before = len(h.log_text())
    h.cell_focus(app)
    h.type_text("64717")
    h.pump(2)
    new_log = h.log_text()[before:]

    check("duplicate warning appears in the log",
          "DUPLICATE" in new_log and "64717" in new_log,
          [l.strip() for l in new_log.splitlines() if "DUPLICATE" in l])
    check("warning names both locations",
          "Co-tau" in new_log and "Fe-tau" in new_log)
    check("duplicates readout > 0", app.readouts["dupes"].cget("text") == "1",
          app.readouts["dupes"].cget("text"))

    key = app._order()[0]
    check("offending cell is outlined red",
          app.cells[key].cget("highlightbackground") == rsxd_logger.RED,
          app.cells[key].cget("highlightbackground"))
    check("duplicate detected ACROSS peak and temperature",
          set(app.session.duplicates().keys()) == {64717})
    check("no dialog was raised (test would have thrown)", True)

    # And it clears when the typo is fixed.
    h.clear_focused()
    h.type_text("64718")
    h.pump(2)
    check("warning clears once corrected", not app.session.duplicates(),
          str(app.session.duplicates()))
    h.stop()


# =============================================================================
# TEST 5  --  extending the range renames, leaves no stale file
# =============================================================================
def test_5():
    banner(5, "Add a scan that extends the range")
    h = Harness(scratch("t5"))
    app = h.start()
    out = h.workdir / "output"

    h.cell_focus(app)
    for scan in (64717, 64718, 64719):
        h.type_text(str(scan))
        h.key("Tab")
    app.save_current()
    h.pump(2)
    first = sorted(p.name for p in out.glob("*.xlsx"))
    check("initial filename carries the range",
          first == ["CoTau_300K_64717-64719.xlsx"], str(first))

    # extend upward
    h.type_text("64768")
    h.key("Tab")
    app.save_current()
    h.pump(2)
    second = sorted(p.name for p in out.glob("*.xlsx"))
    check("filename updated to the new range",
          second == ["CoTau_300K_64717-64768.xlsx"], str(second))
    check("no stale file left behind", len(second) == 1, f"{len(second)} file(s)")
    check("log records the rename", "renamed from CoTau_300K_64717-64719.xlsx"
          in h.log_text())

    # extend downward too
    h.type_text("64700")
    h.key("Tab")
    app.save_current()
    h.pump(2)
    third = sorted(p.name for p in out.glob("*.xlsx"))
    check("extending downward also renames, still one file",
          third == ["CoTau_300K_64700-64768.xlsx"], str(third))

    # a second peak/temperature must NOT be swept up by the cleanup
    app.peak_var.set("Fe-tau"); app.on_context_change(); h.pump(2)
    h.cell_focus(app)
    h.type_text("65000"); h.key("Tab")
    app.save_all()
    h.pump(2)
    allf = sorted(p.name for p in out.glob("*.xlsx"))
    check("other peak's workbook untouched",
          allf == ["CoTau_300K_64700-64768.xlsx", "FeTau_300K_65000-65000.xlsx"],
          str(allf))
    h.stop()


# =============================================================================
# TEST 6  --  present vs absent .dat look different, neither pops up
# =============================================================================
def test_6():
    banner(6, "One scan whose .dat is present, one whose .dat is absent")
    if not Path(REAL_DAT_DIR).is_dir():
        check("real .dat folder available", False, REAL_DAT_DIR)
        return
    h = Harness(scratch("t6"))
    app = h.start()
    app.dat_var.set(REAL_DAT_DIR)
    app.apply_dirs()
    h.pump(2)
    check(".dat check enabled against the real folder", app.session.checker.enabled,
          REAL_DAT_DIR)

    h.cell_focus(app)
    h.type_text(str(REAL_DAT_LO))          # exists on disk
    h.key("Tab")
    h.type_text("99999")                   # does not
    h.pump(2)

    k_present, k_absent = app._order()[0], app._order()[1]
    e_present = app.session.get(app.session.key(*k_present))
    e_absent = app.session.get(app.session.key(*k_absent))
    check("present scan marked found", e_present.dat_found == "yes", e_present.dat_found)
    check("absent scan marked not found", e_absent.dat_found == "no", e_absent.dat_found)

    c_present = app.cells[k_present].cget("highlightbackground")
    c_absent = app.cells[k_absent].cget("highlightbackground")
    check("the two cells are visually different", c_present != c_absent,
          f"present={c_present}  absent={c_absent}")
    check("present = green", c_present == rsxd_logger.GREEN, c_present)
    check("absent = amber (not red, not an error)", c_absent == rsxd_logger.AMBER,
          c_absent)
    check("neither entry raised a dialog (test would have thrown)", True)
    check("filename pattern used is the real i10-1-<scan>.dat form",
          app.session.checker.check(REAL_DAT_LO) == "yes"
          and (Path(REAL_DAT_DIR) / f"i10-1-{REAL_DAT_LO}.dat").is_file())

    # verify all
    before = len(h.log_text())
    app.verify_all()
    h.pump(2)
    vlog = h.log_text()[before:]
    check("verify all reports the missing one",
          "1 still missing" in vlog and "99999" in vlog,
          [l.strip() for l in vlog.splitlines() if "missing" in l])
    h.stop()


# =============================================================================
# TEST 7  --  invalid data directory disables the check silently
# =============================================================================
def test_7():
    banner(7, "Point the data directory somewhere invalid")
    h = Harness(scratch("t7"))
    app = h.start()
    bogus = r"Z:\this\does\not\exist\at\all"
    app.dat_var.set(bogus)
    app.apply_dirs()
    h.pump(2)

    check("checker reports itself disabled", not app.session.checker.enabled)
    check("no exception, app still alive", app.winfo_exists() == 1)
    check("header says DISABLED", "DISABLED" in app.dat_state.cget("text"),
          repr(app.dat_state.cget("text")))

    h.cell_focus(app)
    h.type_text("64717")
    h.key("Tab")
    h.type_text("64718")
    h.pump(2)
    scans = app.session.scans_for(app.session.ui_peak, app.session.ui_temp)
    check("logging still works with the check off", scans == [64717, 64718], str(scans))

    e = app.session.get(app.session.key(*app._order()[0]))
    check("dat_found recorded as 'unknown'", e.dat_found == "unknown", e.dat_found)
    # "logged but unchecked" is the ABSENCE of information, so it is neutral
    # grey. The accent is reserved for focus and would read as a fourth status.
    check("cells use the neutral 'check disabled' colour",
          app.cells[app._order()[0]].cget("highlightbackground") == rsxd_logger.T3,
          app.cells[app._order()[0]].cget("highlightbackground"))
    check("missing readout shows a dash, not a scary number",
          app.readouts["missing"].cget("text") == "—",
          app.readouts["missing"].cget("text"))

    app.save_current()
    h.pump(2)
    files = sorted(p.name for p in (h.workdir / "output").glob("*.xlsx"))
    check("saving still works", files == ["CoTau_300K_64717-64718.xlsx"], str(files))

    before = len(h.log_text())
    app.verify_all()
    h.pump(2)
    check("verify all says so calmly, no crash",
          "check stays disabled" in h.log_text()[before:])

    # and it recovers when pointed somewhere real
    if Path(REAL_DAT_DIR).is_dir():
        app.dat_var.set(REAL_DAT_DIR)
        app.apply_dirs()
        h.pump(2)
        check("recovers when pointed at a real folder", app.session.checker.enabled)
    h.stop()


# =============================================================================
# TEST 8  --  exported .xlsx matches the template layout
# =============================================================================
TEMPLATE = r"G:\My Drive\PhD_Hexaferrites_and_Transport\99_Misc\Claude_staging\Data structure.xlsx"


def test_8():
    banner(8, "Exported .xlsx matches the template layout")
    h = Harness(scratch("t8"))
    app = h.start()
    h.cell_focus(app)
    scan = 64717
    for _ in range(52):
        h.type_text(str(scan))
        h.key("Tab")
        scan += 1
    # a note, to prove notes stay out of the grid cells
    h.cell_focus(app, app._order()[5])
    app.note_var.set("beam dump, repeated")
    h.pump(1)
    app.save_current()
    h.pump(2)

    path = next((h.workdir / "output").glob("*.xlsx"))
    ws = openpyxl.load_workbook(path)["Sheet1"]

    if Path(TEMPLATE).is_file():
        tws = openpyxl.load_workbook(TEMPLATE)["Sheet1"]
        check("A2 header identical to template",
              ws["A2"].value == tws["A2"].value, repr(ws["A2"].value))
        angles_t = [tws.cell(row=2, column=c).value for c in range(3, 16)]
        angles_o = [ws.cell(row=2, column=c).value for c in range(3, 16)]
        check("angle header row identical to template",
              angles_o == angles_t, str(angles_o))
        rows_t = [tws.cell(row=r, column=2).value for r in range(3, 7)]
        rows_o = [ws.cell(row=r, column=2).value for r in range(3, 7)]
        check("polarisation row labels and order identical to template",
              rows_o == rows_t == ["p", "n", "LH", "LV"], str(rows_o))
        merges_t = {str(m) for m in tws.merged_cells.ranges}
        merges_o = {str(m) for m in ws.merged_cells.ranges}
        check("group-label merges identical to template",
              merges_o == merges_t == {"A3:A4", "A5:A6"}, str(sorted(merges_o)))
        check("column widths A-H carried over from template",
              all(abs((ws.column_dimensions[c].width or 0)
                      - (tws.column_dimensions[c].width or 0)) < 1e-6
                  for c in "ABCDEFGH"))
        check("template's own labels were the reversed ones (as the brief said)",
              tws["A3"].value == "Linear dichroism"
              and tws["A5"].value == "Circular Dichroism",
              f"template A3={tws['A3'].value!r} A5={tws['A5'].value!r}")
    else:
        check("template file reachable", False, TEMPLATE)

    check("CORRECTED grouping: p/n = circular",
          ws["A3"].value == "Circular dichroism", repr(ws["A3"].value))
    check("CORRECTED grouping: LH/LV = linear",
          ws["A5"].value == "Linear dichroism", repr(ws["A5"].value))

    grid = [[ws.cell(row=r, column=c).value for c in range(3, 16)] for r in range(3, 7)]
    check("all 52 scan numbers present in the grid",
          sorted(v for row in grid for v in row) == list(range(64717, 64769)))
    check("scan cells are integers, not text",
          all(isinstance(v, int) for row in grid for v in row))
    check("no cell contains a note", not any(isinstance(v, str) for row in grid for v in row))

    wb = openpyxl.load_workbook(path)
    check("notes live on a separate sheet", wb.sheetnames == ["Sheet1", "Notes"],
          str(wb.sheetnames))
    nws = wb["Notes"]
    # The Notes sheet now leads with the gain-epoch summary, so the data table
    # starts wherever that ends. Find the header row rather than hard-coding a
    # row number that will move again the next time something is added.
    head_row = next(r for r in range(1, 40)
                    if nws.cell(row=r, column=1).value == "analyser_angle_deg")
    cols = [nws.cell(row=head_row, column=c).value for c in range(1, 7)]
    check("notes sheet carries angle, polarisation, scan, gain, note, dat_found",
          cols == ["analyser_angle_deg", "polarisation", "scan_number", "gain",
                   "note", "dat_found"], str(cols))
    body = [[nws.cell(row=r, column=c).value for c in range(1, 7)]
            for r in range(head_row + 1, head_row + 60)
            if nws.cell(row=r, column=3).value is not None]
    noted = [row for row in body if row[4]]
    check("note exported alongside its scan number",
          len(noted) == 1 and noted[0][4] == "beam dump, repeated"
          and noted[0][2] == 64722, str(noted[:1]))
    check("every logged scan appears with a gain column, not just noted ones",
          len(body) == 52 and all(row[3] for row in body),
          f"{len(body)} rows")
    check("gain reads 'unknown' when none was recorded, never 1",
          {row[3] for row in body} == {"unknown"},
          str(sorted({row[3] for row in body})))
    check("gain epoch summary is on the sheet",
          nws["A5"].value == "Gain epochs (V/A)", repr(nws["A5"].value))
    h.stop()


# =============================================================================
# TEST 9  --  cumulative CSV holds everything, long format
# =============================================================================
def test_9():
    banner(9, "Cumulative CSV contains every scan, across all peaks and temperatures")
    h = Harness(scratch("t9"))
    app = h.start()

    expected = {}
    scan = 65100
    for peak, temp in (("Co-tau", "300"), ("Co-tau", "200"), ("Fe2", "10")):
        app.peak_var.set(peak)
        app.temp_var.set(temp)
        app.on_context_change()
        h.pump(2)
        h.cell_focus(app)
        for i in range(5):
            focused = h.root.focus_get()
            cell = next(k for k, w in app.cells.items() if w is focused)
            h.type_text(str(scan))
            expected[(peak, temp, cell[0], cell[1])] = scan
            scan += 1
            h.key("Tab")
    app.flush()
    h.pump(2)

    import csv as _csv
    csv_path = h.workdir / "output" / "rsxd_scan_log.csv"
    check("CSV exists at the expected path", csv_path.is_file(), str(csv_path))
    rows = list(_csv.DictReader(csv_path.open(encoding="utf-8")))

    check("column set is exactly the frozen schema",
          list(rows[0].keys()) == CSV_COLUMNS, str(list(rows[0].keys())))
    check("one row per scan, 15 rows", len(rows) == 15, str(len(rows)))

    got = {(r["peak"], r["temperature_K"], r["analyser_angle_deg"], r["polarisation"]):
           int(r["scan_number"]) for r in rows}
    check("every scan present with the right peak/temp/angle/polarisation",
          got == expected, f"{len(got)} rows matched")
    check("all three peak/temperature pairs represented",
          {(r["peak"], r["temperature_K"]) for r in rows}
          == {("Co-tau", "300"), ("Co-tau", "200"), ("Fe2", "10")})
    check("long format: 4 polarisation channels appear as separate rows",
          {r["polarisation"] for r in rows} >= {"p", "n", "LH", "LV"},
          str(sorted({r["polarisation"] for r in rows})))
    check("dat_found column populated", {r["dat_found"] for r in rows} <= {"yes", "no", "unknown"})
    check("timestamp column populated", all(r["timestamp"] for r in rows))

    # correcting a typo must not leave the wrong value behind in the CSV
    h.cell_focus(app, app._order()[0])
    h.clear_focused()
    h.type_text("77777")
    app.flush()
    h.pump(2)
    rows2 = list(_csv.DictReader(csv_path.open(encoding="utf-8")))
    nums = {int(r["scan_number"]) for r in rows2}
    check("corrected typo replaced, old value gone",
          77777 in nums and 65110 not in nums and len(rows2) == 15,
          f"{len(rows2)} rows")
    h.stop()


# =============================================================================
# TEST 10  --  change the temperature list in the config, restart
# =============================================================================
def test_10():
    banner(10, "Change the temperature list in the config, restart")
    work = scratch("t10")
    h = Harness(work)
    app = h.start()
    before = list(app.cfg.temperatures)
    check("default temperature list loaded from config",
          before == ["300", "280", "260", "240", "220", "200", "180", "160",
                     "140", "120", "100", "80", "60", "40", "20", "10"], str(before))
    check("selector offers exactly that list",
          list(app.temp_box.cget("values")) == before)

    h.cell_focus(app)
    h.type_text("64717")
    app.flush()
    h.stop()

    # Edit the plain-text config by hand, exactly as the user would.
    text = h.config_path.read_text(encoding="utf-8")
    text = re.sub(r"^temperatures_K\s*=.*$",
                  "temperatures_K = [295, 150, 77, 4.2]", text, count=1, flags=re.M)
    h.config_path.write_text(text, encoding="utf-8")

    h2 = Harness(work, reuse_config=True)
    app2 = h2.start()
    after = list(app2.cfg.temperatures)
    check("restart picks up the new temperature list",
          after == ["295", "150", "77", "4.2"], str(after))
    check("selector now offers the new list",
          list(app2.temp_box.cget("values")) == ["295", "150", "77", "4.2"])
    check("non-integer temperature survives (4.2 K)", "4.2" in after)
    check("app fell back to a valid temperature", app2.session.ui_temp in after,
          app2.session.ui_temp)
    check("earlier 300 K scan not silently dropped",
          any(k[1] == "300" for k in app2.session.entries))
    check("and the user is told about it",
          "no longer match the config" in h2.log_text(),
          [l.strip() for l in h2.log_text().splitlines() if "no longer match" in l])

    # a new workbook at 4.2 K, to prove the whole chain works on the new list
    app2.temp_var.set("4.2"); app2.on_context_change(); h2.pump(2)
    h2.cell_focus(app2)
    h2.type_text("70001")
    app2.save_current(); h2.pump(2)
    files = sorted(p.name for p in (work / "output").glob("*4.2K*.xlsx"))
    check("workbook written for the new 4.2 K point",
          files == ["CoTau_4.2K_70001-70001.xlsx"], str(files))

    # angles and peaks are equally not hardcoded
    text = h2.config_path.read_text(encoding="utf-8")
    text = re.sub(r"^analyser_angles_deg\s*=.*$",
                  "analyser_angles_deg = [-30, 0, 30]", text, count=1, flags=re.M)
    text = re.sub(r'^"Co-tau" = .*$', '"Co-tau" = "CobaltTau"', text, count=1, flags=re.M)
    h2.config_path.write_text(text, encoding="utf-8")
    h2.stop()

    h3 = Harness(work, reuse_config=True)
    app3 = h3.start()
    check("angle list is config-driven too", app3.cfg.angles == ["-30", "0", "30"],
          str(app3.cfg.angles))
    check("grid rebuilt with 3 columns x 4 rows", len(app3.cells) == 12,
          str(len(app3.cells)))
    check("peak filename string is config-driven",
          app3.cfg.peak_file_string("Co-tau") == "CobaltTau")
    h3.stop()


# =============================================================================
# TEST 11  --  sequential fill (column / row / block), always editable after
# =============================================================================
def test_11():
    banner(11, "Optional sequential fill: row, column and block")
    h = Harness(scratch("t11"))
    app = h.start()
    cfg = app.cfg

    # ---- column ----
    h.cell_focus(app, (cfg.angles[0], cfg.pol_labels[0]))
    app.fill_start.delete(0, "end"); app.fill_start.insert(0, "64717")
    app.fill_mode.set("column")
    app.do_fill(); h.pump(2)
    col = [app.session.get(app.session.key(cfg.angles[0], p)) for p in cfg.pol_labels]
    check("column fill wrote 4 cells at one angle, incrementing",
          [e.scan for e in col] == [64717, 64718, 64719, 64720],
          str([e.scan for e in col]))
    check("fill did not touch the next angle",
          app.session.get(app.session.key(cfg.angles[1], "p")) is None)

    # ---- still editable afterwards (real measurement order differs) ----
    h.cell_focus(app, (cfg.angles[0], "LH"))
    h.clear_focused()
    h.type_text("64999")
    h.pump(2)
    check("a filled cell stays editable afterwards",
          app.session.get(app.session.key(cfg.angles[0], "LH")).scan == 64999)

    # ---- row ----
    h.cell_focus(app, (cfg.angles[0], "p"))
    app.fill_start.delete(0, "end"); app.fill_start.insert(0, "65200")
    app.fill_step.delete(0, "end"); app.fill_step.insert(0, "2")
    app.fill_mode.set("row")
    app.do_fill(); h.pump(2)
    scans = [(app.session.get(app.session.key(a, "p")) or Entry0).scan
             for a in cfg.angles]
    check("row fill left the already-filled cell alone", scans[0] == 64717,
          str(scans[:3]))
    check("row fill walked the remaining 12 angles with step 2",
          scans[1:] == list(range(65200, 65200 + 12 * 2, 2)), str(scans[1:]))

    # ---- block, on a clean grid ----
    h2 = Harness(scratch("t11b"))
    app2 = h2.start()
    h2.cell_focus(app2)
    app2.fill_start.delete(0, "end"); app2.fill_start.insert(0, "70000")
    app2.fill_mode.set("block")
    app2.do_fill(); h2.pump(2)
    filled = app2.session.scans_for(app2.session.ui_peak, app2.session.ui_temp)
    check("block fill wrote the whole 52-cell grid",
          filled == list(range(70000, 70052)), f"{len(filled)} cells")
    check("block order is column-major (4 channels, then next angle)",
          app2.session.get(app2.session.key(cfg.angles[0], "LV")).scan == 70003
          and app2.session.get(app2.session.key(cfg.angles[1], "p")).scan == 70004)
    check("no duplicates produced", not app2.session.duplicates())
    check("fill is reported in the log, not silent",
          "Sequential fill: 52 cells" in h2.log_text(),
          [l.strip() for l in h2.log_text().splitlines() if "Sequential fill" in l])
    h2.stop()
    h.stop()


# =============================================================================
# TEST 12  --  per-cell notes: keyboard reachable, exported separately
# =============================================================================
def test_12():
    banner(12, "Per-cell note field")
    h = Harness(scratch("t12"))
    app = h.start()
    key = h.cell_focus(app)
    h.type_text("64717")
    h.pump(2)

    check("note bar binds to the focused cell",
          "-90" in app.note_label.cget("text") and "p" in app.note_label.cget("text"),
          app.note_label.cget("text"))

    h.key("Control-n")
    h.pump(2)
    check("Ctrl+N moves focus to the note field, no mouse",
          h.root.focus_get() is app.note_entry)

    for ch in "beam":
        app.note_entry.event_generate("<KeyPress>", keysym=ch)
    h.pump(2)
    check("typing in the note field writes through to the cell",
          app.session.get(app.session.key(*key)).note == "beam",
          repr(app.session.get(app.session.key(*key)).note))

    h.key("Return")
    h.pump(2)
    check("Return returns focus to the grid cell", h.root.focus_get() is app.cells[key])
    # Has-note is real state, not decoration, so it gets its own colour --
    # not the focus accent (which would collide with the focused cell) and not
    # one of the three status hues (which mean something else).
    check("cell number turns violet to show it has a note",
          app.cells[key].cget("fg") == rsxd_logger.VIOLET,
          app.cells[key].cget("fg"))
    check("the has-note marker is distinct from the focus accent",
          rsxd_logger.VIOLET not in (rsxd_logger.CYAN, rsxd_logger.AMBER,
                                     rsxd_logger.GREEN, rsxd_logger.RED))
    check("the scan number itself is untouched by the note",
          app.cell_vars[key].get() == "64717")

    app.save_current(); h.pump(2)
    import csv as _csv
    csv_path = h.workdir / "output" / "rsxd_scan_log.csv"
    rows = list(_csv.DictReader(csv_path.open(encoding="utf-8")))
    check("note is its own CSV column, never inside scan_number",
          rows[0]["note"] == "beam" and rows[0]["scan_number"] == "64717",
          str(rows[0]))

    # a note on an empty cell survives, without inventing a scan number
    k2 = app._order()[1]
    h.cell_focus(app, k2)
    app.note_var.set("bad alignment, no scan taken")
    h.pump(2)
    app.flush(); h.pump(2)
    rows = list(_csv.DictReader(csv_path.open(encoding="utf-8")))
    check("note on an EMPTY cell does not fabricate a CSV row",
          len(rows) == 1, f"{len(rows)} rows")

    # clearing a mistyped scan number must not take the note with it
    h.cell_focus(app, key)
    h.clear_focused()
    h.pump(2)
    check("clearing the scan number keeps the cell's note",
          app.session.get(app.session.key(*key)).note == "beam",
          repr(app.session.get(app.session.key(*key))))
    check("but the cleared scan is gone from the CSV",
          not app.session.scans_for(app.session.ui_peak, app.session.ui_temp))
    h.type_text("64720")
    h.pump(2)
    check("retyping a number reunites it with the surviving note",
          app.session.get(app.session.key(*key)).scan == 64720
          and app.session.get(app.session.key(*key)).note == "beam")
    h.stop()

    h2 = Harness(h.workdir, reuse_config=True)
    app2 = h2.start()
    check("note on an empty cell still survives a restart",
          any(e.note == "bad alignment, no scan taken"
              for e in app2.session.entries.values()))
    h2.stop()


# =============================================================================
# TEST 13  --  starts from one command; embeds as a tab without a rewrite
# =============================================================================
def test_13():
    banner(13, "Single-command launch, and embedding as a tab")
    work = scratch("t13")
    Harness(work)
    cfg_path = work / "rsxd_config.toml"

    def run(name, body):
        p = work / name
        p.write_text(f"import sys\nsys.path.insert(0, r\"{APP_DIR}\")\n"
                     f"CFG = r\"{cfg_path}\"\n" + body, encoding="utf-8")
        return subprocess.run([sys.executable, str(p)], capture_output=True,
                              text=True, timeout=120)

    # ---- launches and survives a real mainloop ----
    proc = run("launch.py", (
        "import tkinter as tk, rsxd_logger\n"
        "root = tk.Tk()\n"
        "app = rsxd_logger.RSXDLoggerFrame(root, config_path=CFG)\n"
        "app.pack(fill='both', expand=True)\n"
        "app.install_shortcuts(root)\n"
        "w, h = app.preferred_size()\n"
        "root.geometry('%dx%d' % (w, h))\n"
        "root.after(900, lambda: (app.on_close(), root.destroy()))\n"
        "root.mainloop()\n"
        "print('LAUNCH_OK', w, h, flush=True)\n"
    ))
    check("app starts from a single command and exits cleanly",
          "LAUNCH_OK" in proc.stdout and proc.returncode == 0,
          (proc.stdout.strip() + " " + proc.stderr.strip()[-300:]).strip())

    bat = APP_DIR / "run_rsxd_logger.bat"
    check("double-click launcher exists", bat.is_file(), str(bat))

    # ---- embeds inside a ttk.Notebook as one tab among several ----
    proc = run("embed.py", (
        "import tkinter as tk\n"
        "from tkinter import ttk\n"
        "import rsxd_logger\n"
        "root = tk.Tk()\n"
        "nb = ttk.Notebook(root); nb.pack(fill='both', expand=True)\n"
        "other = ttk.Frame(nb); nb.add(other, text='Something Else')\n"
        "page = ttk.Frame(nb); nb.add(page, text='Scan Log')\n"
        "app = rsxd_logger.RSXDLoggerFrame(page, config_path=CFG)\n"
        "app.pack(fill='both', expand=True)\n"
        "app.install_shortcuts(root)\n"
        "for _ in range(8): root.update_idletasks(); root.update()\n"
        "nb.select(page)\n"
        "for _ in range(6): root.update_idletasks(); root.update()\n"
        "assert len(app.cells) == 52, len(app.cells)\n"
        "k = app._order()[0]\n"
        "app.cells[k].focus_force()\n"
        "for _ in range(3): root.update_idletasks(); root.update()\n"
        "for ch in '64717': app.cells[k].event_generate('<KeyPress>', keysym=ch)\n"
        "for _ in range(3): root.update_idletasks(); root.update()\n"
        "assert app.session.get(app.session.key(*k)).scan == 64717\n"
        "app.on_close()\n"
        "print('EMBED_OK', len(nb.tabs()), flush=True)\n"
        "root.destroy()\n"
    ))
    check("embeds as a ttk.Notebook tab, no rewrite, still logs scans",
          "EMBED_OK 2" in proc.stdout and proc.returncode == 0,
          (proc.stdout.strip() + " " + proc.stderr.strip()[-400:]).strip())

    # ---- the core is importable with no GUI at all (phase 2 reuse) ----
    proc = run("core_only.py", (
        "import builtins\n"
        "real = builtins.__import__\n"
        "def guard(name, *a, **k):\n"
        "    if name.split('.')[0] == 'tkinter':\n"
        "        raise AssertionError('rsxd_core pulled in tkinter')\n"
        "    return real(name, *a, **k)\n"
        "builtins.__import__ = guard\n"
        "import rsxd_core\n"
        "cfg = rsxd_core.Config.load(CFG)\n"
        "s = rsxd_core.Session(cfg); s.load_session()\n"
        "print('CORE_OK', len(cfg.angles), len(cfg.pol_labels), flush=True)\n"
    ))
    check("rsxd_core works with tkinter forbidden — phase 2 can reuse it headless",
          "CORE_OK 13 4" in proc.stdout and proc.returncode == 0,
          (proc.stdout.strip() + " " + proc.stderr.strip()[-400:]).strip())


# =============================================================================
def main():
    tests = [test_1, test_2, test_3, test_4, test_5,
             test_6, test_7, test_8, test_9, test_10,
             test_11, test_12, test_13]
    only = sys.argv[1:]
    for t in tests:
        if only and t.__name__.split("_")[-1] not in only:
            continue
        try:
            t()
        except Exception as exc:
            import traceback
            traceback.print_exc()
            RESULTS.append((f"{t.__name__} CRASHED", False, str(exc)))

    print(f"\n{'=' * 78}")
    passed = sum(1 for _, ok, _ in RESULTS if ok)
    print(f"SUMMARY   {passed}/{len(RESULTS)} checks passed")
    for name, ok, detail in RESULTS:
        if not ok:
            print(f"  FAIL  {name}   {detail}")
    print("=" * 78)
    return 0 if passed == len(RESULTS) else 1


if __name__ == "__main__":
    raise SystemExit(main())
